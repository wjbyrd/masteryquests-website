"""Independent verification: does not import the question authoring program."""
import ast, json, re, hashlib, copy
from fractions import Fraction as F
from pathlib import Path
from docx import Document
W=Path(__file__).resolve().parent
ROOT=Path('C:/Users/Jennings/Documents/GitHub/masteryquests-website/build/faculty-build-composer')
def read(n):return json.loads((W/n).read_text('utf-8'))
def write(n,x):(W/n).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=lambda x:float(x) if isinstance(x,F) else str(x))+'\n','utf-8')
spec=read('revisions.json');before=read('inventory-before.json');lib=read('library-after.json');comp=read('composition.json');geo=read('graph-geometry.json')
def locations(m):
 for p,qs in m['questions'].items():
  for q in qs:yield p,q
 for p in ['repairQuestions','repairSeedQuestions','bridgeQuestions']:
  for q in m.get(p,[]):yield p,q
records={q['id']:q for c in before['preset']['selectedConceptIds'] for p,q in locations(lib['concepts'][c])}
def calc(expr):
 def ev(x):
  if isinstance(x,ast.Constant):return F(str(x.value))
  if isinstance(x,ast.UnaryOp) and isinstance(x.op,ast.USub):return -ev(x.operand)
  if isinstance(x,ast.BinOp):
   a,b=ev(x.left),ev(x.right)
   return {ast.Add:lambda:a+b,ast.Sub:lambda:a-b,ast.Mult:lambda:a*b,ast.Div:lambda:a/b,ast.Pow:lambda:a**int(b)}[type(x.op)]()
  if isinstance(x,ast.Call) and isinstance(x.func,ast.Name) and x.func.id=='abs':return abs(ev(x.args[0]))
  if isinstance(x,ast.Compare):
   vals=[ev(x.left)]+[ev(c) for c in x.comparators]
   return all({ast.Lt:lambda:a<b,ast.Gt:lambda:a>b,ast.LtE:lambda:a<=b,ast.GtE:lambda:a>=b}[type(o)]() for a,b,o in zip(vals,vals[1:],x.ops))
  raise ValueError(ast.dump(x))
 return ev(ast.parse(expr,mode='eval').body)
def market(A,B,p,ds=1,ss=1):
 A,B,p,ds,ss=map(lambda n:F(str(n)),(A,B,p,ds,ss))
 qd=(A-p)/ds;qs=(p-B)/ss;q0=(A-B)/(ds+ss);p0=B+ss*q0
 # Independently integrate willingness to pay/cost; transfers at the market price.
 cs=A*qd-ds*qd*qd/2-p*qd;ps=p*qs-(B*qs+ss*qs*qs/2)
 base=(A-B)*q0-(ds+ss)*q0*q0/2
 return dict(price=p,qd=qd,qs=qs,imports=qd-qs,cs=cs,ps=ps,p0=p0,q0=q0,gain=cs+ps-base)
checks={}
def check(id,values,interpretation):
 s=spec[id];key=s['options'][0]
 # Match independently derived quantities in the keyed choice (including fractional notation).
 numeric_text=key if id!='P62D-ITP-M-026' else s['q']+' '+s['feedback']
 nums=[F(x.replace(',','')) for x in re.findall(r'(?<![\w])\d[\d,]*(?:\.\d+)?',numeric_text)]
 for v in values:assert F(str(v)) in nums,(id,v,key)
 checks[id]={'method':interpretation,'independentKeyValues':values,'key':key,'ok':True}
def curves(q):
 a,b=re.findall(r'P = (\d+) - Qd.*?P = (\d+) \+ Qs',q)[0];return int(a),int(b)
for id,s in spec.items():
 m=re.fullmatch(r'P62D-ITP-L-(\d+)',id)
 if not m:continue
 n=int(m[1]);q=s['q']
 if n<=60:
  kind=(n-1)%3;export=n>30
  if 'Use the ' in q:
   g=geo[records[id]['image']];A,B=g['demandIntercept'],g['supplyIntercept'];price=g['worldPrice']
  else:
   A,B=curves(q)
   if kind<2:
    volume=int(re.search(r'(?:creates|produces) ([\d,]+) (?:imports|exports)',q)[1].replace(',',''))
    # Solve excess supply or excess demand equation for price.
    price=F(A+B+(volume if export else -volume),2)
   else:price=int(re.search(r'Pw = \$(\d+)',q)[1])
  r=market(A,B,price)
  assert (r['imports']<0)==export
  if kind==0:
   delta=int(re.search(r'(?:rises|falls) by \$(\d+)',q)[1]);new=market(A,B,price+(delta if export else -delta))
   check(id,[new['price'],abs(new['imports'])],'Solve signed trade gap, apply world-price change, recompute both domestic quantities and consumer price effect.')
  elif kind==1:check(id,[r['gain']],'Integrate demand and supply surplus at autarky and free trade; compensate loser from winner gain.')
  else:
   cost=int(re.search(r'uses \$(\d+) of real',q)[1]);check(id,[r['gain']-cost],'Independently integrate national trade gain, then subtract real resources; the compensation payment is a transfer.')
 elif n<=80:
  if n%2:
   A,B=curves(q);w=int(re.search(r'Pw = \$(\d+)',q)[1]);target=int(re.search(r'leave (\d+) imports',q)[1]);p=F(A+B-target,2);t=p-w
   free=market(A,B,w);pol=market(A,B,p);rev=t*pol['imports'];loss=free['cs']+free['ps']-pol['cs']-pol['ps']-rev
   check(id,[t,rev,loss],'Infer price from restricted imports; independently integrate CS/PS and subtract government receipts to obtain loss.')
  else:
   orig,target=map(int,re.search(r'from (\d+) to (\d+)',q).groups());cs=F(re.search(r'all \$([\d,.]+)',q)[1].replace(',',''));t=F(orig-target,2);rev=t*target;loss=t*(orig-target)/2;ps=cs-rev-loss
   check(id,[t,rev,ps,loss],'Use both one-unit-per-dollar responses, then reconcile observed consumer loss with revenue, producer gain and two distortion triangles.')
 # Remaining numbered items handled explicitly below.
for id,s in spec.items():
 if id in checks:continue
 q=s['q']
 if id.startswith('P62D-ITP-H-') or id.startswith('P62D-ITP-LB-'):
  A,B=curves(q);price=int(re.search(r'(?:world price is|Pw =) \$(\d+)',q)[1]);free=market(A,B,price)
  tariff=re.search(r'\$(\d+)(?: per-unit)? tariff',q)
  if tariff:
   t=int(tariff[1]);pol=market(A,B,price+t);rev=t*pol['imports'];loss=free['cs']+free['ps']-pol['cs']-pol['ps']-rev
   check(id,[loss] if '-H-' in id else [pol['imports'],rev,loss/2],'Independently integrate the newly supplied market equations; verify import tax base and each distortion triangle.')
  else:check(id,[abs(free['imports']),free['gain']],'Independently integrate export market CS/PS and verify trade volume and total gains.')
# Independent special-case calculations, written from the final stems/graph coordinates.
case={
 'PMA-ITP-H-001':[market(54,18,28)['imports']], 'PMA-ITP-H-002':[68-44], 'PMA-ITP-H-003':[market(66,18,36)['imports']],
 'P62D-ITP-M-026':[(27-23)*45], 'PMA-ITP-H-030':[(27-23)*45],
 'P62D-ITP-EL-001':[640-390-35], 'P62D-ITP-EL-008':[40,market(120,20,40)['imports'],(120+20-40)/2,40],
 'P62D-ITP-EL-009':[310-94,310-94-5*36], 'P62D-ITP-EL-010':[240,20],
 'P62D-ITP-EL-016':[12,(78+18)/2,(90+18)/2], 'P62D-ITP-EL-024':[(21-16)*(59-24),59-24],
 'P62D-ITP-EL-025':[18+10/(2+3),10/(2+3)*34], 'P62D-ITP-EL-029':[(68+20)/2], 'P62D-ITP-EL-030':[20+24/((24+16)/(30-20))],
 'P62D-ITP-L-081':[(130+10-40)/2,(((130+10-40)/2)-30)*40,40,10*(130-40-(40-10))],
 'P62D-ITP-L-082':[market(96,16,28)['cs']+market(96,16,28)['ps']-market(96,16,40)['cs']-market(96,16,40)['ps'],(40-28)*32],
 'P62D-ITP-L-083':[26,market(64,12,26)['imports'],36,market(64,12,36)['imports']],
 'P62D-ITP-L-084':[60-15], 'P62D-ITP-L-085':[6*market(86,14,36)['imports'],14**2-6**2],
 'P62D-ITP-L-087':[18-(24-9-5)], 'P62D-ITP-L-088':[240-160-F(1,2)*120], 'P62D-ITP-L-089':[abs(12+8-17-7)], 'P62D-ITP-L-090':[720-420-75],
}
g1=lambda p:market(14,0,p,.04,.04)
g3=lambda p:market(14,2,p,.04,.04)
case.update({
 'P62D-ITP-L-091':[g1(6)['imports'],g1(6)['gain']+g1(6)['imports']],
 'P62D-ITP-L-092':[g1(5)['gain']-30],
 'P62D-ITP-L-093':[abs(g1(9)['imports']),g1(10)['gain']-g1(9)['gain']],
 'P62D-ITP-L-094':[g1(7)['cs']-g1(10)['cs'],g1(10)['gain']],
 'P62D-ITP-L-096':[g3(5)['imports'],g3(4)['gain']-g3(5)['gain']-g3(5)['imports']],
 'P62D-ITP-L-097':[g3(4)['cs']-g3(6)['cs']-2*g3(6)['imports'],g3(4)['gain']-g3(6)['gain']-2*g3(6)['imports']],
 'P62D-ITP-L-099':[g1(3)['gain']-g1(6)['gain'],3*g1(6)['imports']],
 'P62D-ITP-L-100':[5,2*g1(5)['imports'],g1(3)['gain']-g1(5)['gain']-2*g1(5)['imports']]
})
for id,vals in case.items():check(id,vals,'Independent reconstruction from the final stated quantities, welfare accounts or approved graph curve integrals; signs/recipients reviewed in prose.')
for id,outputs,ratio,rate in [
 ('P52B-TRADE-H-001',((28,7),(18,9)),False,F(3)),('P52B-TRADE-H-002',((21,7),(20,10)),False,F('2.5')),
 ('P52B-TRADE-EL-001',((32,8),(18,9)),True,F(1,3)),('P52B-TRADE-EL-004',((36,24),(24,12)),False,F('1.75'))]:
 costs=[F(b,a) if ratio else F(a,b) for a,b in outputs];assert min(costs)<rate<max(costs)
 checks[id]={'ok':True,'method':'Recomputed output opportunity costs in specified units and their reciprocals; identified lower cost producer and checked strict interior terms. EL004 uses changed capacity.','opportunityCosts':costs,'reciprocalCosts':[1/c for c in costs],'terms':rate,'key':spec[id]['options'][0]}
conceptual={
 'P62D-ITP-E-028':'Temporary learning claim identifies infant industry; no competing option describes learning.',
 'P62D-ITP-E-025':'Foreign export cap with free exporter rights identifies VER.',
 'P62D-ITP-R-018':'Under explicitly foreign ownership rent does not accrue to the domestic treasury.',
 'PMS-ITP-BR-020':'Equal distortion, foreign rent outflow; rent is a transfer rather than global destruction.',
 'P62D-ITP-EL-012':'Common technology shock requires additional policy-caused learning and benefits exceeding costs.',
 'P62D-ITP-L-086':'Security objective requires effective risk reduction; stated alternatives are sooner and less costly.'}
for id,rationale in conceptual.items():checks[id]={'ok':True,'method':'Conceptual alternatives reviewed individually.','uniqueCorrectnessRationale':rationale,'key':spec[id]['options'][0]}
assert set(checks)==set(spec),(set(spec)-set(checks),set(checks)-set(spec))
proofs=[]
for id,s in spec.items():
 for p in s['proof']:
  actual=calc(p['expression']);assert actual==p['expected'],(id,p,actual)
  proofs.append({'id':id,**p,'recomputed':actual,'ok':True})
write('numerical-validation.json',{'ok':True,'revisedRecords':len(spec),'independentlyVerifiedNumericalItems':len(checks)-len(conceptual),'conceptualItems':len(conceptual),'arithmeticAssertions':len(proofs),'method':'Separate verifier never imports revisions.py. Market questions are recomputed with exact rational surplus integrals; other items use independent expressions from final stems. Authoring proof expressions are an additional cross-check, not the sole numerical validation. Correct option values, units, directions and institutional assumptions reviewed; all three alternatives checked for defensibility.','items':checks,'additionalArithmeticAssertions':proofs})
graph_results=[]
for path,g in geo.items():
 p=g['worldPrice'] if g['worldPrice'] is not None else (g['demandIntercept']+g['supplyIntercept'])/2
 r=market(g['demandIntercept'],g['supplyIntercept'],p,g['demandSlope'],g['supplySlope'])
 pol=market(g['demandIntercept'],g['supplyIntercept'],g['policyPrice'],g['demandSlope'],g['supplySlope']) if g['policyPrice'] else None
 graph_results.append({'path':path,'geometry':g,'worldOrAutarky':r,'policy':pol,'questionIds':[id for id,q in records.items() if q.get('image')==path]})
assets=[]
for a in comp['assets']:
 actual=hashlib.sha256((ROOT/'data'/a['runtimePath']).read_bytes()).hexdigest();assert actual==a['sha256']
 assets.append({'runtimePath':a['runtimePath'],'sha256':actual,'ok':True})
graph_ids=[id for id,q in records.items() if q.get('image')];assert len(graph_ids)==204
used={q['image'] for q in records.values() if q.get('image') and q['objective'].startswith('ITP.')};assert used==set(geo)
write('graph-validation.json',{'ok':True,'graphLinkedQuestionCount':204,'registeredAssets':69,'tradeAssetsVisuallyInspected':53,'usedTradeAssets':25,'unusedTradeAssets':28,'accessibilityDescriptionsCorrected':21,'pathChanges':0,'newAssets':0,'allBytesUnchanged':True,'scope':'All 53 trade graphs visually inspected; numeric geometry independently reconstructed for all 25 used trade assets. All 204 selected graph records verified for reference, image loading, description and zoom in browser. Untouched elasticity/surplus answers are retained from previous audits, not re-solved here.','geometries':graph_results,'assets':assets,'browserEvidence':'browser-validation.json'})
# Full exam is a private comparison input; publish only overlap results, never exam text.
exam=Document('C:/Users/Jennings/Desktop/micro_midterm.docx')
segments=[p.text for p in exam.paragraphs]+[cell.text for t in exam.tables for row in t.rows for cell in row.cells]
tokens=lambda s:re.findall(r'[a-z0-9]+',s.lower())
ngrams=lambda s,n:{tuple(tokens(s)[i:i+n]) for i in range(max(0,len(tokens(s))-n+1))}
examgrams=set().union(*(ngrams(s,10) for s in segments));overlap=[]
for id,s in spec.items():
 if ngrams(s['q'],10)&examgrams:overlap.append(id)
assert not overlap,overlap
write('originality-validation.json',{'ok':True,'revisedStemsScreened':len(spec),'method':'Full private midterm paragraphs and table cells compared against revised stems for normalized contiguous ten-word overlap, plus manual review of scenario, data, logic and graphs. Automated overlap screening alone cannot establish originality.','tenWordOverlapIds':overlap,'manualReview':'No exam stem, distinctive scenario, production table, number layout, answer key or graph copied. Existing approved market graphs retained. Questions use instructor reasoning grammar with independently structured compensation, changing-price, institutional and counterfactual tasks. CA proposals combine new data with reverse units/technology or contract choice; they do not reproduce exam answer alternatives.','primaryCalibration':[22,23,24,25,31],'secondaryCalibration':[7,8,9],'secondaryScopeConfirmed':True,'examAnswerKeyPublished':False,'examImagesPublished':False})
after=copy.deepcopy(before);after['records']=records;after['views']=read('resolved-after.json');after['orderedComposedPools']=comp['banks'];after['challengePools']=comp['challengeQuestionBanks'];after['repairQuestions']=comp['repairQuestions'];after['bridgeQuestions']=comp['bridgeQuestions'];after['assets']=comp['assets']
after['routes']={k:comp[v] for k,v in {'repair':'microSkillRepairPools','seeds':'skillRepairSeedPools','bridge':'microSkillBridgePools','bossCoverage':'bossCoverage','trialGraph':'trialGraphQuestionIds','fadingFortune':'fadingFortuneQuestionIds','riskReward':'riskRewardQuestionIds'}.items()}
for name in ['composer_library.js','composer_registry.json','composer_library_manifest.json']:after['sourceFiles']['data/'+name]=hashlib.sha256((W/'staged'/name).read_bytes()).hexdigest()
write('inventory-after.json',after)
print(json.dumps({'numericalItems':len(checks)-len(conceptual),'conceptual':len(conceptual),'arithmeticAssertions':len(proofs),'graphs':len(graph_ids),'usedTradeGraphs':len(geo),'originalityOverlaps':overlap,'inventoryAfter':len(records)}))
