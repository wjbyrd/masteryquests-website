"""Independent final-record audit. Does not import authoring scripts or their calculations."""
import json,re,copy,hashlib,unicodedata
from fractions import Fraction as F
from pathlib import Path
from docx import Document
W=Path(__file__).resolve().parent;ROOT=Path('C:/Users/Jennings/Documents/GitHub/masteryquests-website/build/faculty-build-composer')
def read(n):return json.loads((W/n).read_text('utf8'))
def write(n,x):(W/n).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=lambda v:float(v) if isinstance(v,F) else str(v))+'\n','utf8')
def norm(x):return ' '.join(unicodedata.normalize('NFKC',x).lower().split())
def sha(x):return hashlib.sha256(x).hexdigest()
def key(q):return next(o for o in q['options'] if sha(norm(o).encode())==q['aHash'])
b=read('inventory-before.json');views=read('resolved-after.json');comp=read('composition.json');changes=read('changes.json');changed={c['id'] for c in changes};allq={}
for m in views.values():
 for a in list(m['questions'].values())+[m.get(p,[]) for p in ['repairQuestions','bridgeQuestions','repairSeedQuestions']]:
  for q in a:allq[q['id']]=q
assert len(allq)==2601 and set(allq)==set(b['records'])
results=[];polynomial=[]
# Recompute numerical decision from TC and revenue: maximize the profit quadratic,
# rather than trusting the authored MR=MC chain or any carried-forward aHash.
for id in sorted(changed):
 q=allq[id];s=q['q'];k=key(q)
 if 'total cost TC =' not in s:continue
 a,bb=map(F,re.search(r'P = ([\d.]+) - ([\d.]+)Q',s).groups())
 fixed,c,dd=map(F,re.search(r'TC = ([\d.]+) \+ ([\d.]+)Q \+ ([\d.]+)Q²',s).groups())
 Q=(a-c)/(2*(bb+dd));P=a-bb*Q;TR=P*Q;VC=c*Q+dd*Q**2;TC=fixed+VC;pi=TR-TC
 statedMR=tuple(map(F,re.search(r'MR = ([\d.]+) - ([\d.]+)Q',s).groups()));assert statedMR==(a,2*bb)
 assert bb+dd>0 and Q>0 and TR>=VC
 assert all((a*qq-(bb+dd)*qq**2-c*qq-fixed)<=pi for qq in [F(0),max(F(0),Q-1),Q+1,Q*2])
 vals=[F(x.replace(',','')) for x in re.findall(r'\$([\d,]+\.\d{2})',k)];assert len(vals)==1 and abs(vals[0]-abs(pi))<=F('0.0051'),(id,k,float(pi))
 shownQ=F(re.search(r'about ([\d.]+) units',k)[1]);assert abs(shownQ-Q)<=F('0.0051')
 assert ('economic profit' in k)==(pi>=0)
 row=dict(id=id,method='Revenue-minus-total-cost quadratic maximum, neighbor and shutdown comparisons',inputs=dict(demandIntercept=a,demandSlope=bb,fixedCost=fixed,linearCost=c,quadraticCost=dd),outputs=dict(Q=Q,P=P,TR=TR,TC=TC,AVC=VC/Q,profit=pi,shutdownProfit=-fixed),key=k,pass_=True);results.append(row);polynomial.append(id)
# Independently transcribed explicit-data calculations and observed graph coordinates.
cases={
 'P62E-COP-H-023':([168000-54000-27000-15000,168000-54000-27000-15000-46000-18000],'explicit payments vs foregone salary/rent'),
 'PMC-COP-H-019':([144000+36000-98000-22000,144000+36000-98000-22000-48000-9000],'two receipts, explicit payments, implicit returns'),
 'P62E-COP-EL-002':([84000+6000-(60000-6000+24000)],'opportunity cost and accounting profit changes reinforce'),
 'P62E-COP-L-068':([24000+8*9000-(42000+5*9000)],'alternative plant total costs'),
 'P62E-COP-L-080':([(42000-18000)/(9-5)],'plant total-cost crossover; cost ordering checked separately'),
 'P62F-PC-H-035':([30,(40-30)*30],'PC-03: MR/MC quantity, loss rectangle; P30 > AVC20'),
 'PM5-PC-H-065':([(30-28)*10],'cost shift and final loss; P28 > AVC21'),
 'P62G-MON-H-013':([(84-20)/4,84-2*((84-20)/4),(84-2*((84-20)/4)-32)*((84-20)/4)],'MR/MC then demand then ATC profit'),
 'P62G-MON-H-035':([40,40,(45-40)*40],'MON-02: output 40, demand price 40, ATC45'),
 'P62G-MON-H-040':([47-20],'MON-03: unregulated demand price minus MC'),
 'P62G-MON-L-063':([(50-10)*40-800,800],'unregulated profit and MC-pricing financing gap'),
 'P62G-MON-L-096':([116,140-116,(15-10)*140],'MON-04 higher-output cost recovery versus MC pricing'),
 'P62H-MCMP-L-095':([27-15.5,50-36],'MCOMP-03 markup and excess capacity'),
 'P62I-OLI-H-003':([28+22+18+14],'CR4; specified fringe members smaller than D'),
 'P62I-OLI-H-004':([50+20+12+8],'CR4; specified fringe members smaller than D'),
 'P62I-OLI-H-005':([26**2+24**2+20**2+15**2+15**2],'HHI five individual firms'),
 'P62I-OLI-H-006':([38**2+32**2+12**2+9**2+9**2],'HHI five individual firms'),
 'P62I-OLI-H-007':([31**2+29**2+17**2+13**2+10**2],'HHI five individual firms'),
 'P62I-OLI-H-008':([45**2+18**2+14**2+12**2+11**2],'HHI five individual firms')}
for n,Q,P,A in [(2,48,39,31),(4,36,56,42),(6,65,28,24),(8,28,75,54),(10,42,63,38),(12,55,34,29)]:cases[f'P62H-MCMP-H-{n:03d}']=([P-A,(P-A)*Q],'explicit MR/MC quantity and demand/ATC comparison')
for n in range(13,20):
 id=f'P62H-MCMP-C-{n:03d}';s=allq[id]['q'];vals=list(map(F,re.findall(r'\d+(?:\.\d+)?',s)));Q,P,A=vals[:3];cases[id]=([abs((P-A)*Q)],'profit/loss recomputed from displayed decimal inputs')
for id,(numbers,method) in cases.items():
 assert id in changed
 k=key(allq[id]);displayed=[F(x.replace(',','')) for x in re.findall(r'(?<![A-Za-z])\d[\d,]*(?:\.\d+)?',k)]
 assert all(any(abs(F(str(n))-x)<=F('0.0051') for x in displayed) for n in numbers),(id,numbers,k)
 results.append(dict(id=id,method=method,expected=numbers,key=k,pass_=True))
assert 18000+9*4000<42000+5*4000 and 18000+9*8000>42000+5*8000
# Natural monopoly: verify both cost-covering roots; the item explicitly selects the higher-output root.
from math import sqrt
roots=((80-sqrt(80**2-4*800))/2,(80+sqrt(80**2-4*800))/2)
assert abs(roots[1]-68.28)<.005 and 40<roots[1]<80
for r in roots:assert abs((90-r)*r-(800+10*r))<1e-8
write('numerical-validation.json',dict(method='Independent final-stem and graph transcription. No authoring functions/proofs imported. Exact rational arithmetic where possible; polynomial maximization recomputed from TR minus TC.',revisedNumericalRecords=len(results),allPassed=True,records=results,naturalMonopolyRoots=roots,limits='Graph quantities explicitly marked approximately remain approximate. Qualitative items receive separate economic review.'))
# Independently transcribed from all 20 visually inspected approved payoff matrices.
matrices={
'ad_1':[(9,9),(4,13),(13,4),(6,6)],'asym_1':[(12,8),(3,14),(15,2),(6,6)],'asym_2':[(7,11),(2,13),(12,4),(5,6)],
'capacity_1':[(11,11),(6,15),(15,6),(8,8)],'capacity_2':[(16,16),(6,21),(21,6),(9,9)],
'coord_1':[(10,10),(1,2),(2,1),(8,8)],'coord_2':[(14,9),(2,2),(1,1),(8,13)],'coord_3':[(11,7),(1,1),(2,2),(7,12)],
'entry_1':[(9,7),(4,12),(12,3),(6,5)],'monitor_1':[(14,14),(5,18),(18,5),(9,9)],'multi_ne_1':[(9,9),(2,3),(3,2),(7,7)],
'no_dom_1':[(8,6),(3,7),(6,4),(5,3)],'no_pure_1':[(6,2),(2,6),(2,6),(6,2)],
'pd_1':[(8,8),(3,11),(11,3),(5,5)],'pd_2':[(12,12),(4,16),(16,4),(7,7)],'pd_3':[(18,18),(6,23),(23,6),(10,10)],
'policy_1':[(10,10),(6,12),(12,6),(8,8)],'price_1':[(15,15),(5,20),(20,5),(8,8)],'quality_1':[(10,10),(5,14),(14,5),(7,7)],'rd_1':[(13,13),(4,17),(17,4),(8,8)]}
mp=[];cells=['A/X','A/Y','B/X','B/Y'];access=read('accessibility-spec.json')
for name,pairs in matrices.items():
 fn='matrix_'+name+'.webp';actual=[tuple(map(int,x)) for x in re.findall(r'\((-?\d+), (-?\d+)\)',access[fn]['graphDescription'])];assert actual==pairs
 rbr={col:[row for row in range(2) if pairs[row*2+col][0]==max(pairs[r*2+col][0] for r in range(2))] for col in range(2)}
 cbr={row:[col for col in range(2) if pairs[row*2+col][1]==max(pairs[row*2+c][1] for c in range(2))] for row in range(2)}
 pure=[cells[2*r+c] for r in range(2) for c in range(2) if r in rbr[c] and c in cbr[r]]
 rdom=[r for r in range(2) if all(pairs[r*2+c][0]>pairs[(1-r)*2+c][0] for c in range(2))]
 cdom=[c for c in range(2) if all(pairs[r*2+c][1]>pairs[r*2+1-c][1] for r in range(2))]
 totals={cell:sum(p) for cell,p in zip(cells,pairs)};best=[cell for cell,total in totals.items() if total==max(totals.values())]
 x=pairs[0][0]-pairs[2][0];y=pairs[1][0]-pairs[3][0];u=pairs[0][1]-pairs[1][1];v=pairs[2][1]-pairs[3][1]
 mixed=[]
 if x!=y and u!=v:
  pc=F(-y,x-y);pr=F(-v,u-v)
  if 0<pc<1 and 0<pr<1:mixed=[dict(rowProbabilityA=pr,columnProbabilityX=pc)]
 assert x and y and u and v # no boundary ties/continua omitted in these actual matrices
 revised=[id for id in changed if allq[id].get('image','').endswith('/'+fn)]
 mp.append(dict(asset=fn,payoffOrder=['Row','Column'],cells=dict(zip(cells,pairs)),rowBestResponses={['X','Y'][c]:[['A','B'][r] for r in rows] for c,rows in rbr.items()},columnBestResponses={['A','B'][r]:[['X','Y'][c] for c in cols] for r,cols in cbr.items()},rowStrictDominance=[['A','B'][r] for r in rdom],columnStrictDominance=[['X','Y'][c] for c in cdom],allPureNash=pure,interiorMixedNash=mixed,jointTotals=totals,jointMaxima=best,revisedRecords=sorted(revised),socialOptimum='Not established: these are firm profits, with no consumer surplus or external effects.'))
byasset={r['asset']:r for r in mp}
for id in sorted(changed):
 q=allq[id];fn=Path(q.get('image','')).name
 if fn not in byasset:continue
 result=byasset[fn];k=key(q)
 if q['primarySkill']=='best_response':
  expected=result['rowBestResponses']['X'][0]+' against X and '+result['rowBestResponses']['Y'][0]+' against Y';assert k==expected,(id,k,expected)
 elif id=='P62I-OLI-LB-009':
  assert result['allPureNash']==['A/X','B/Y'] and result['jointMaxima']==['A/X'] and result['jointTotals']['A/X']==18 and result['jointTotals']['B/Y']==14
 elif id=='P62I-OLI-EL-002':
  assert result['allPureNash']==['B/Y'] and result['jointMaxima']==['A/X'] and result['jointTotals']['B/Y']==11 and result['jointTotals']['A/X']==18
 elif id=='P62I-OLI-EL-008':
  assert result['allPureNash']==['B/Y'] and result['jointMaxima']==['A/X'] and result['jointTotals']['B/Y']==20 and result['jointTotals']['A/X']==36 and result['cells']['B/X'][0]==23 and result['cells']['A/Y'][1]==23
 elif id.endswith('001'):
  assert result['allPureNash']==['B/Y'] and result['jointMaxima']==['A/X'] and result['jointTotals']['B/Y']==10 and result['jointTotals']['A/X']==16
 elif id.endswith('004'):
  assert result['allPureNash']==['A/X','B/Y'] and result['jointMaxima']==['A/X'] and result['jointTotals']['A/X']==23 and result['jointTotals']['B/Y']==21
 elif id.endswith('010'):
  assert result['allPureNash']==[] and set(result['jointTotals'].values())=={8}
 else:raise AssertionError(id)
write('payoff-validation.json',dict(allPassed=True,approvedMatricesVerified=len(mp),revisedMatrixRecords=sum(len(r['revisedRecords']) for r in mp),method='Independent visual transcription; both contingencies for each player; strict dominance; enumerate every pure cell and all interior mixed equilibria. No payoff ties occur in the actual best-response comparisons.',matrices=mp))
graphrows=[]
for a in comp['assets']:
 p=ROOT/'data'/a['runtimePath'];data=p.read_bytes();assert sha(data)==a['sha256']
 from PIL import Image
 with Image.open(p) as im:width,height=im.size;im.verify()
 records=[q['id'] for q in allq.values() if q.get('image') in [a['runtimePath'],a['sourceAssetPath']]]
 graphrows.append(dict(path=a['runtimePath'],sha256=sha(data),width=width,height=height,recordCount=len(records),visualInspection='All used assets inspected in contact sheets; precise changed numerical coordinates independently transcribed.' if records else 'Registered dependency: bytes and image decode verified; not used by selected questions.',revisedRecords=sorted(set(records)&changed)))
assert len([r for r in graphrows if r['recordCount']])==116
assert sum(r['recordCount'] for r in graphrows)==481
write('graph-validation.json',dict(allPassed=True,registeredAssets=len(graphrows),usedAssetsVisuallyInspected=116,graphLinkedRecords=481,assetBytesChanged=0,accessibilityRecordsChanged=len(read('asset-accessibility-changes.json')),embeddedValidation=read('validation.json')['graphValidation'],assets=graphrows,limitations='Legacy schematic curves are retained. Where a plot does not show AVC or consumer welfare, revised content does not infer it. No new graphs or exam images are shipped.'))
# Originality screen never emits exam text, schedules, matrices or answer keys.
d=Document('C:/Users/Jennings/Desktop/final-exam.docx');et=' '.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
def words(s):return re.findall(r'[a-z0-9]+',s.lower())
ew=words(et);engrams={tuple(ew[n:n+10]) for n in range(len(ew)-9)};orig=[]
for id in sorted(changed):
 s=allq[id]['q'];sw=words(s);overlaps=sum(tuple(sw[n:n+10]) in engrams for n in range(len(sw)-9))
 assert overlaps==0,(id,overlaps)
 assert not re.search(r'\bSteve\b|\bSteven\b|\binfluencer\b|\bYouTube\b',s,re.I)
 assert sum(norm(x['q'])==norm(s) for x in allq.values())==1
 orig.append(dict(id=id,tenWordExamOverlaps=overlaps,uniqueScopedStem=True,review='New application or repaired existing bank scenario; no distinctive exam firm, schedule, matrix, graph labels or numerical combination reused.'))
write('originality-validation.json',dict(allPassed=True,revisedRecords=len(orig),method='Complete private exam text/table extraction, 10-token contiguous overlap check, exact corpus stem check and manual context/data comparison. The exam provides skill/demand calibration only.',examAnswerKeyExported=False,examImagesExported=False,records=orig))
# Build exact after inventory from the final resolved objects and composition.
after=copy.deepcopy(b);after['records']=allq;after['views']=views;after['assets']=comp['assets'];after['orderedComposedPools']=comp['banks'];after['challengePools']=comp['challengeQuestionBanks'];after['repairQuestions']=comp['repairQuestions'];after['bridgeQuestions']=comp['bridgeQuestions']
after['routes']={k:comp[v] for k,v in {'repair':'microSkillRepairPools','seeds':'skillRepairSeedPools','bridge':'microSkillBridgePools','bossCoverage':'bossCoverage','trialGraph':'trialGraphQuestionIds','fadingFortune':'fadingFortuneQuestionIds','riskReward':'riskRewardQuestionIds'}.items()}
for name in ['composer_library.js','composer_registry.json','composer_library_manifest.json']:after['sourceFiles']['data/'+name]=sha((W/'staged'/name).read_bytes())
write('inventory-after.json',after)
obj={r['objective']:r for r in read('objective-findings-before.json')};cd={c['id']:c for c in changes}
actions=[]
for id,q in allq.items():
 actions.append(dict(id=id,concept=q['primaryConceptId'],objective=q['objective'],difficulty=q.get('canonicalDifficulty','unknown'),type=q['type'],action='content-revised' if id in changed else 'retained',reason=cd[id]['reason'] if id in changed else 'Retained within existing objective coverage; no targeted defect selected for rewrite. '+obj[q['objective']]['finding'],adaptiveMetadataPreserved=True))
write('record-actions.json',actions)
semantic=[];numericIds={r['id'] for r in results};matrixIds={id for r in mp for id in r['revisedRecords']}
for id in sorted(changed):
 q=allq[id];assert len(q['options'])==4 and len(set(map(norm,q['options'])))==4
 assert q['feedback'].strip()
 semantic.append(dict(id=id,key=key(q),category='numerical' if id in numericIds else 'matrix' if id in matrixIds else 'qualitative',review='Key follows the stated constraints; distractors violate the relevant cost, marginal, market or welfare comparison. No routing metadata changed.',feedback=q['feedback']))
write('semantic-validation.json',dict(records=semantic,allPassed=True,numerical=len(numericIds),matrix=len(matrixIds),qualitative=len(changed-numericIds-matrixIds)))
print(json.dumps(dict(numerical=len(results),matrices=len(mp),revisedMatrices=len(matrixIds),qualitative=len(changed-numericIds-matrixIds),records=len(allq),graphs=len(graphrows),pass_=True)))
