"""Build a game-specific guide from its editable Markdown source.

Use the bundled Codex Python runtime (python-docx), from the repository root.
Economy's Edge uses the Faculty Game Overview as its reference; Signal House
uses the final Economy's Edge guide for page geometry, branding and styles.
"""
from copy import deepcopy
from pathlib import Path
import argparse
from io import BytesIO
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser()
parser.add_argument('--game',choices=['the-economys-edge','signal-house','gdp-live','cpi-live'],default='the-economys-edge')
game=parser.parse_args().game
title={'signal-house':'Signal House','gdp-live':'GDP Live','cpi-live':'CPI Live','the-economys-edge':'The Economy’s Edge'}[game]
SOURCE = ROOT / f'docs/faculty-guides/{game}.md'
OUTPUT = ROOT / f'downloads/resources/{game}-faculty-guide.docx'
REFERENCE = ROOT / ('downloads/resources/signal-house-faculty-guide.docx' if game=='gdp-live' else 'downloads/resources/the-economys-edge-faculty-guide.docx' if game=='signal-house' else 'downloads/resources/faculty-game-overview-guide.docx')
if game=='cpi-live':
    REFERENCE=ROOT/'downloads/resources/gdp-live-faculty-guide.docx'

doc = Document(REFERENCE)
logo = deepcopy(next(p._p for p in doc.paragraphs if p._p.xpath('.//w:drawing')))
for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'):
        doc._element.body.remove(child)
doc._element.body.insert(0, logo)
# Keep branding but remove the reference game's unused teaching image.
logo_images={node.get(qn('r:embed')) for node in logo.xpath('.//a:blip')}
for rel_id,rel in list(doc.part.rels.items()):
    if rel.reltype==RT.IMAGE and rel_id not in logo_images:
        doc.part.drop_rel(rel_id)
doc.core_properties.title = title+' Faculty Guide'
doc.core_properties.subject = 'Teaching economic mechanisms through evidence and aggregate supply' if game=='signal-house' else 'Teaching production possibilities through a branching economics game'
doc.core_properties.author = 'Mastery Quests'
doc.core_properties.keywords = 'Faculty guide, aggregate supply, AD-AS, evidence, policy tradeoffs' if game=='signal-house' else 'Faculty guide, production possibilities frontier, PPF, opportunity cost'
doc.core_properties.comments = ''
if game=='gdp-live':
    doc.core_properties.subject='Teaching GDP expenditure accounting through transaction classification and ledger correction'
    doc.core_properties.keywords='Faculty guide, GDP, expenditure identity, imports, investment, national accounts'
if game=='cpi-live':
    doc.core_properties.subject='Teaching fixed-basket price measurement, CPI, and inflation interpretation'
    doc.core_properties.keywords='Faculty guide, CPI, fixed basket, inflation, expenditure weights, measurement'
for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'List Bullet', 'List Number']:
    style = doc.styles[style_name]
    style.font.name = 'Aptos'
    style.font.color.rgb = RGBColor.from_string('172B45')
    style.paragraph_format.widow_control = True
for style_name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[style_name].font.color.rgb = RGBColor.from_string('000000')
    doc.styles[style_name].paragraph_format.keep_with_next = True
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Subtitle'].font.color.rgb = RGBColor(0,0,0)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(9)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].paragraph_format.space_after = Pt(8)
for name in ['List Bullet', 'List Number']:
    doc.styles[name].font.size = Pt(10.5)
    doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.styles[name].paragraph_format.line_spacing = 1.08
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(13)
for section in doc.sections:
    for p in section.header.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0,0,0)
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run('masteryquests.org  •  '+title+'  •  ')
    r.font.name = 'Aptos'; r.font.size = Pt(8)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'en-US')
if doc.styles['Normal'].element.get_or_add_rPr().find(qn('w:lang')) is None:
    doc.styles['Normal'].element.get_or_add_rPr().append(lang)

def text(paragraph, value):
    for token in re.split(r'(\*\*.*?\*\*|\[.*?\]\(.*?\))', value):
        if token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif re.fullmatch(r'\[.*?\]\(.*?\)', token):
            label, url = re.fullmatch(r'\[(.*?)\]\((.*?)\)', token).groups()
            rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
            link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
            run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '0F345E'); props.append(color)
            underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
            run.append(props); content = OxmlElement('w:t'); content.text = label; run.append(content)
            link.append(run); paragraph._p.append(link)
        else:
            paragraph.add_run(token)

def numbered_list(bullet=False):
    numbering = doc.part.numbering_part.element
    abstract_id = max(int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))) + 1
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), str(abstract_id))
    nsid=OxmlElement('w:nsid'); nsid.set(qn('w:val'),f'{abstract_id+7000:08X}'); abstract.append(nsid)
    multi=OxmlElement('w:multiLevelType'); multi.set(qn('w:val'),'singleLevel'); abstract.append(multi)
    level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
    for name, val in [('start','1'),('numFmt','bullet' if bullet else 'decimal'),('lvlText','•' if bullet else '%1.'),('lvlJc','left')]:
        element=OxmlElement('w:'+name); element.set(qn('w:val'),val); level.append(element)
    props=OxmlElement('w:pPr'); indent=OxmlElement('w:ind'); indent.set(qn('w:left'),'360'); indent.set(qn('w:hanging'),'360'); props.append(indent);level.append(props)
    abstract.append(level)
    numbering.insert(list(numbering).index(numbering.find(qn('w:num'))),abstract)
    num=numbering.add_num(abstract_id)
    num.add_lvlOverride(0).add_startOverride(1)
    return num.numId

bullet_id=numbered_list(bullet=True)

lines=SOURCE.read_text(encoding='utf-8').splitlines()
i=0; num_id=None
while i < len(lines):
    line=lines[i].strip(); i+=1
    if not line:
        continue
    if line == '<!-- pagebreak -->':
        doc.add_page_break(); continue
    if line.startswith('# '):
        p=doc.add_paragraph('GAME-SPECIFIC FACULTY GUIDE', style='Eyebrow')
        p=doc.add_paragraph(line[2:],style='Title'); continue
    if line.startswith('## '):
        doc.add_paragraph(line[3:],style='Heading 1'); num_id=None; continue
    if line.startswith('### '):
        doc.add_paragraph(line[4:],style='Heading 2'); num_id=None; continue
    if line.startswith('!['):
        alt, image_path = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line).groups()
        # Signal House's preceding paragraph covers hints, not the figure.
        if game!='signal-house':
            doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        image_source=(SOURCE.parent/image_path).resolve()
        if image_source.suffix=='.webp':
            from PIL import Image
            converted=BytesIO()
            Image.open(image_source).save(converted,format='PNG')
            converted.seek(0);image_source=converted
        else:image_source=str(image_source)
        shape=p.add_run().add_picture(image_source,width=Inches(5.25))
        shape._inline.docPr.set('descr', alt); shape._inline.docPr.set('title','Signal House environment' if game=='signal-house' else 'Calder’s four production regions')
        continue
    if line.startswith('|'):
        rows=[line]
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i].strip()); i+=1
        rows=[row for row in rows if not re.fullmatch(r'[| :\-]+',row)]
        table=doc.add_table(rows=0,cols=2); table.autofit=False
        table.columns[0].width=Inches(3.45); table.columns[1].width=Inches(3.55)
        for index,row in enumerate(rows):
            cells=table.add_row().cells
            trpr=table.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if index==0:
                trpr.append(OxmlElement('w:tblHeader'))
            for cell,value in zip(cells,row.strip('|').split('|')):
                p=cell.paragraphs[0]; text(p,value.strip()); p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(6)
                for run in p.runs: run.font.size=Pt(10)
                tcpr=cell._tc.get_or_add_tcPr()
                shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'0F345E' if index==0 else ('F1F5F8' if index%2 else 'FFFFFF'));tcpr.append(shading)
                if index==0:
                    for run in p.runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        continue
    if line.startswith('*') and line.endswith('*') and not line.startswith('**'):
        style='Subtitle' if ('branching PPF experience' in line or 'two-case economic escape room' in line or 'GDP accounting game' in line or 'consumer-price measurement game' in line) else 'Caption'
        p=doc.add_paragraph(line[1:-1],style=style);continue
    if line.startswith('Faculty Resource Guide'):
        continue
    if line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet')
        props=p._p.get_or_add_pPr().get_or_add_numPr();props.get_or_add_ilvl().val=0;props.get_or_add_numId().val=bullet_id
        text(p,line[2:]);num_id=None;continue
    match=re.match(r'^(\d+)\. (.*)$',line)
    if match:
        if match[1]=='1' or num_id is None:num_id=numbered_list()
        p=doc.add_paragraph(style='List Number');props=p._p.get_or_add_pPr().get_or_add_numPr();props.get_or_add_ilvl().val=0;props.get_or_add_numId().val=num_id
        text(p,match[2]);continue
    num_id=None
    p=doc.add_paragraph(); text(p,line)
# Keep the short optional question set together without a fixed page break.
if game in ['signal-house','gdp-live','cpi-live']:
    in_concepts=False
    for paragraph in doc.paragraphs:
        if paragraph.style.name=='Heading 1':
            in_concepts=paragraph.text=='What the Game Is Really Teaching'
        elif in_concepts:
            paragraph.paragraph_format.keep_together=True
question_group=[]
in_questions=False
for paragraph in doc.paragraphs:
    if paragraph.style.name == 'Heading 1':
        in_questions = paragraph.text == 'Questions to Consider'
    if in_questions:
        question_group.append(paragraph)
for paragraph in question_group[:-1]:
    paragraph.paragraph_format.keep_with_next=True

doc.save(OUTPUT)
print(OUTPUT)
