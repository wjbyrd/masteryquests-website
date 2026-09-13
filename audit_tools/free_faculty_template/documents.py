from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from lxml import html
import json

root=Path(__file__).resolve().parents[2]
resources=root/'downloads/resources'
evidence=root/'validation_artifacts/free_faculty_template_parity/documents'
evidence.mkdir(parents=True,exist_ok=True)
names=['faculty-composer-quick-start-guide.docx','faculty-game-overview-guide.docx','faculty-implementation-guide.docx','student-instructions-and-faculty-customization-checklist.docx','faculty-question-bank-helper.docx','example-question-architecture.docx','example-question-generation-prompt.docx']
for name in names:
    if not (evidence/name).exists(): copy2(resources/name,evidence/name)

def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs(cell)

privacy='Public polished games and the free manual template keep progress and activity data in the browser. Remote anonymous telemetry is OFF by default. The Composer offers an explicit optional operational telemetry build; operational telemetry is not research by default. Application telemetry does not intentionally include names, email addresses, LMS account identifiers, university identifiers or free-text responses. Learners choose whether to share local downloads through the course workflow.'
for name in names[:4]:
    doc=Document(evidence/name)
    for p in paragraphs(doc):
        if 'Separate classroom research builds' in p.text:
            p.text=privacy
        elif 'Progress and run data stay in local browser storage and are not automatically sent to Mastery Quests.' in p.text:
            p.text=p.text.replace('Progress and run data stay in local browser storage and are not automatically sent to Mastery Quests.','Remote collection is OFF by default; optional operational builds require explicit faculty configuration.')
        elif p.text.startswith('Current online guidance:'):
            p.text='Current guidance: https://masteryquests.org/how-to/composer/ | Learning outcomes: https://masteryquests.org/how-to/learning-outcomes/ | Manual authoring and validation: https://masteryquests.org/downloads/resources/manual-faculty-guide.html'
    doc.save(resources/name)

def new_doc(title):
    doc=Document()
    sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor.from_string('111111')
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Title'].font.size=Pt(24)
    doc.add_paragraph(title,'Title')
    return doc

soup=html.fromstring((resources/'manual-faculty-guide.html').read_text(encoding='utf-8'))
doc=new_doc('Manual faculty question bank guide')
for element in soup.xpath('//main')[0]:
    name=element.tag
    if name=='h2':doc.add_heading(element.text_content(),level=1)
    elif name=='p':doc.add_paragraph(element.text_content())
    elif name in ('ul','ol'):
        for i,item in enumerate(element.findall('li'),1):doc.add_paragraph((str(i)+'. ' if name=='ol' else '')+item.text_content(),'Normal' if name=='ol' else 'List Bullet')
    elif name=='pre':
        p=doc.add_paragraph(element.text_content());p.paragraph_format.keep_together=True
        for run in p.runs:run.font.name='Consolas';run.font.size=Pt(9)
    elif name=='table':
        table=doc.add_table(rows=0,cols=2);table.style='Light Shading Accent 1'
        for row in element.xpath('.//tr'):
            cells=table.add_row().cells
            for cell,text in zip(cells,row):cell.text=text.text_content()
doc.add_paragraph('Current downloads and clickable file links: https://masteryquests.org/downloads/resources/manual-faculty-guide.html')
doc.save(resources/names[4])

sample=json.loads((resources/'faculty-question-package-example.json').read_text())
doc=new_doc('Current faculty question architecture')
doc.add_paragraph('Use a JSON package with banks, repairQuestions, bridgeQuestions, objectiveLabels, embeddedQuestionAssets and questionAssetMetadata. Download the working package rather than copying incomplete fragments from this document.')
doc.add_paragraph('https://masteryquests.org/downloads/resources/faculty-question-package-example.json')
doc.add_heading('A complete ordinary question',level=1)
p=doc.add_paragraph(json.dumps(sample['banks']['medium'][0],indent=2))
for run in p.runs:run.font.name='Consolas';run.font.size=Pt(9)
doc.add_paragraph('Place this object in banks.medium. All questions require id or questionId, q, four nonempty string options, a from 0–3 or aHash, tag, type, objective, primarySkill, repairSkill and feedback. Main/boss difficulty must match its pool. Repair and bridge can omit difficulty. Numeric ID ranges, objectiveLabel and bossEligible are not required.')
doc.add_heading('Images and adaptive routes',level=1)
doc.add_paragraph('For an image question add image: "sample-bars.svg" and, for Trial by Graph, graphRequired: true. In questionAssetMetadata["sample-bars.svg"], supply imageAlt and graphDescription. The downloadable example also embeds the SVG as a base64 image data URL under embeddedQuestionAssets, so no image folder is required. Per-question alt fields alone do not supply the maintained graph dialog metadata.')
doc.add_paragraph('Main repairSkill values must match repair and bridge primarySkill values. Keep all IDs unique. Type is any nonempty string label; custom labels are allowed. Current aliases include questionType, outcomeIds[0] or conceptId, skillId, repairSkillId and canonicalDifficulty. BossStage, when provided, is 1, 2, 3, opening, middle or final.')
doc.add_paragraph('Validate the complete package at https://masteryquests.org/tools/question-bank-validator/ and use the manual guide for pool counts, settings, telemetry and the release checklist. The free template is local-only; remote anonymous telemetry is OFF by default.')
doc.save(resources/names[5])

doc=new_doc('Faculty question generation prompt')
doc.add_paragraph('Supply your learning objectives, source materials and requested pool counts with this prompt. Review every generated answer yourself. Do not provide student records or identifiers.')
doc.add_heading('Prompt to copy',level=1)
for text in [
'Generate a Mastery Quests manual question package as valid JSON only, without markdown fences, comments or executable JavaScript. Use the current sample structure at https://masteryquests.org/downloads/resources/faculty-question-package-example.json. The top-level fields are banks, repairQuestions, bridgeQuestions, objectiveLabels, embeddedQuestionAssets and questionAssetMetadata.',
'Use banks easy, medium, hard, elite, legendary, easyBoss, mediumBoss, finalBoss and legendaryBoss. Every question needs a unique id (string or number), q, exactly four plausible nonempty string options, numeric a from 0–3, tag, type, objective, primarySkill, repairSkill and feedback explaining the answer. Difficulty must match the main pool. Boss pools use easy, medium, hard and legendary respectively; never difficulty boss. Repair and bridge may omit difficulty.',
'Use objective IDs supplied by me and map them to readable names in objectiveLabels. Type is a descriptive nonempty string, such as calculation or interpretation. Do not invent a required numeric ID range or Composer-only metadata. Use optional hint, commonError, conceptCluster or secondarySkills only when useful. If a bossStage is supplied, use 1, 2, 3, opening, middle or final.',
'Build a repair and bridge route for each taught skill: the main question repairSkill must match the primarySkill of its repair and bridge questions. Do not make repeated question wording the only difference between difficulty levels. Include multi-step reasoning only when appropriate to the objective, while keeping the four-option answer model.',
'Only reference images I actually provide or explicitly create approved assets for. Image is a relative filename. In questionAssetMetadata under that same key provide imageAlt and graphDescription describing the information needed without depending on color. Trial by Graph questions require graphRequired true. Do not invent missing image files or use remote image URLs. Embedded images go in embeddedQuestionAssets as base64 image data URLs.',
'Minimum launch counts: Standard and Score need 6 each easy/medium/hard, 3 each easyBoss/mediumBoss/finalBoss, plus repair and bridge. Timed, Exam and Unlimited need 6 each easy/medium/hard plus repair and bridge. Quiz needs 5 each easy/medium/hard. Legendary needs 6 legendary and 3 legendaryBoss. Trial by Graph needs 10 eligible graph questions. Fading Fortune and Risk and Reward need 10 main questions. Provide 15 or 20 eligible questions for those longer targets. These are minimums; use the larger bank sizes I request.'
]:
    if text.startswith('Only reference images'):doc.add_page_break()
    doc.add_paragraph(text)
doc.add_heading('Review and use',level=1)
doc.add_paragraph('Validate the JSON at https://masteryquests.org/tools/question-bank-validator/. Fix errors, review warnings, check factual accuracy and answer keys, then download checked JSON and paste only within the manual package markers. Test all enabled modes, adaptive routes, images, save/resume and completion before sharing. The validator cannot certify teaching quality or whether an answer hash matches its choices.')
doc.add_paragraph('The free manual template and public polished games are local-only. Remote telemetry is OFF by default. Operational telemetry is not research by default. Do not include names, email addresses, LMS account identifiers, university identifiers or free-text student responses in the package.')
doc.save(resources/names[6])
# Strip inherited Word title rules and use semantic black titles for all edited guides.
for name in names:
    doc=Document(resources/name)
    for style in doc.styles:
        if style.type==1 and style.name=='Title':
            style.font.color.rgb=RGBColor(0,0,0)
            for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    for p in doc.paragraphs:
        if p.style.name=='Title':
            for border in p._p.xpath('.//w:pBdr'):border.getparent().remove(border)
            for run in p.runs:run.font.color.rgb=RGBColor(0,0,0);run.font.underline=False
    doc.save(resources/name)
print(json.dumps({'updated':names},indent=2))
